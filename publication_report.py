"""Offline publication exports for the step-10 RSA cluster CSV.

The CSV remains the machine-readable interface. All presentation formats share
one factual report, with missing provenance recorded separately from manuscript
text. No model calls, Word installation, or LaTeX installation are required.
"""
import hashlib
import json
import os
from pathlib import Path
import re
import shutil

DRIVE_RESULTS_ROOT = r"G:\My Drive\Results"
HEADERS = ["Cluster", "Region at peak", "Extent\n(voxels)", "Peak Z", "x (mm)", "y (mm)", "z (mm)"]
REMOTE_WARNING = ("Steps 10/15.10: this is neither Windows nor a Google Colab runtime. "
                  "Skipping Word export and Google Drive operations; no Word dependencies "
                  "will be imported. CSV, LaTeX, Results text and provenance remain available.")


def word_export_supported():
    """Use runtime markers, never import an optional package to detect Colab.

    Merely installing google-colab on a remote server must not enable Word.
    Google Drive's Windows path remains separately gated by mirror_results.
    """
    return os.name == "nt" or bool(os.environ.get("COLAB_RELEASE_TAG") or
                                   os.environ.get("COLAB_BACKEND_VERSION"))


def _metadata(path, warnings):
    """Some historical .json sidecars are actually YAML."""
    import yaml
    if not path.is_file():
        warnings.append(f"Metadata unavailable: {path.name}")
        return {}
    try:
        value = yaml.safe_load(path.read_text(encoding="utf-8"))
        if not isinstance(value, dict):
            raise ValueError("expected a mapping")
        return value
    except (OSError, ValueError, yaml.YAMLError) as exc:
        warnings.append(f"Cannot read {path.name}: {exc}")
        return {}


def _number(value):
    return f"{float(value):.2f}".rstrip("0").rstrip(".")


def _region(value):
    import pandas as pd
    name = " ".join(str(value).split()) if pd.notna(value) else "Unknown"
    if name in {"", "Unknown", "No label"}:
        return "Unlabelled"
    if name == "OutOfAtlas":
        return "Outside atlas"
    # Preserve atlas abbreviations and qualifiers; change only explicit laterality.
    match = re.fullmatch(r"(.+?) ([LR])( .+)?", name)
    if match:
        side = "Left" if match[2] == "L" else "Right"
        return f"{side} {match[1]}{match[3] or ''}"
    return name[0].upper() + name[1:]


def build_report(csv_path, corrected_map, *, dataset, specie, rsa_model,
                 report_title=None, min_dist_mm=8.0, max_peaks_per_cluster=3,
                 atlas_name=None, regression_model=None, group_metadata_path=None):
    """Validate the CSV against its image and collect only supported statements."""
    import nibabel as nib
    import numpy as np
    import pandas as pd
    csv_path, corrected_map = Path(csv_path), Path(corrected_map)
    frame = pd.read_csv(csv_path)
    required = ["cluster_id", "subpeak_id", "cluster_size_vox", "subpeak_Z",
                "subpeak_x_mm", "subpeak_y_mm", "subpeak_z_mm", "region"]
    if not set(required).issubset(frame.columns):
        raise ValueError("Publication report requires the step-10 cluster CSV columns")
    numeric = required[:-1]
    if not np.isfinite(frame[numeric].to_numpy(dtype=float)).all():
        raise ValueError("Cluster table contains non-finite numeric values")
    if frame.duplicated(["cluster_id", "subpeak_id"]).any():
        raise ValueError("Cluster table contains duplicate peak IDs")
    for column in ["cluster_id", "subpeak_id", "cluster_size_vox"]:
        values = frame[column].to_numpy(dtype=float)
        if np.any(values < 1) or np.any(values != np.floor(values)):
            raise ValueError(f"Invalid positive integer column: {column}")
    frame = frame.sort_values(["cluster_id", "subpeak_id"])
    image = nib.load(corrected_map)
    data = image.get_fdata()
    from scipy.ndimage import label
    labelled, n_clusters = label(np.isfinite(data) & (data > 0), np.ones((3, 3, 3)))
    sizes = np.bincount(labelled.ravel())[1:]
    clusters = []
    for cid, group in frame.groupby("cluster_id", sort=True):
        cid = int(cid)
        if cid > n_clusters or group.cluster_size_vox.nunique() != 1:
            raise ValueError("CSV cluster IDs/extents do not match corrected map")
        size = int(group.cluster_size_vox.iloc[0])
        if size != sizes[cid - 1]:
            raise ValueError("CSV cluster extent does not match corrected map")
        peaks = []
        for row in group.itertuples():
            xyz = [float(row.subpeak_x_mm), float(row.subpeak_y_mm), float(row.subpeak_z_mm)]
            ijk_float = nib.affines.apply_affine(np.linalg.inv(image.affine), xyz)
            ijk = tuple(np.rint(ijk_float).astype(int))
            if (not np.allclose(ijk_float, ijk, atol=1e-4)
                    or any(i < 0 or i >= image.shape[d] for d, i in enumerate(ijk))
                    or labelled[ijk] != cid
                    or not np.isclose(data[ijk], row.subpeak_Z, rtol=1e-5)):
                raise ValueError("CSV peak coordinates/statistics do not match corrected map")
            peaks.append(dict(subpeak_id=int(row.subpeak_id), z=float(row.subpeak_Z),
                              xyz_mm=xyz, region=_region(row.region), raw_region=str(row.region)))
        peaks.sort(key=lambda p: (-p["z"], p["subpeak_id"]))
        if not np.isclose(peaks[0]["z"], data[labelled == cid].max(), rtol=1e-5):
            raise ValueError("CSV does not include the cluster maximum")
        clusters.append(dict(cluster_id=cid, size_vox=size, peaks=peaks))
    if len(clusters) != n_clusters:
        raise ValueError("CSV and corrected map have different cluster counts")

    warnings = []
    map_stem = corrected_map.name.removesuffix(".nii.gz")
    correction_path = corrected_map.with_name(map_stem + ".json")
    group_stem = re.sub(r"_zt[^_]+_corrected$", "_mean", map_stem)
    group_path = (Path(group_metadata_path) if group_metadata_path is not None
                  else corrected_map.with_name(group_stem + ".json"))
    correction = _metadata(correction_path, warnings)
    group_metadata = _metadata(group_path, warnings)
    n_voxels = int(sum(c["size_vox"] for c in clusters))
    # Refuse contradictory receipts instead of writing a polished false report.
    for key, expected in [("n_clusters", n_clusters), ("n_voxels", n_voxels),
                          ("corrected_map", corrected_map.name)]:
        if key in correction and correction[key] != expected:
            raise ValueError(f"Correction metadata disagrees with image: {key}")
    for key, expected in [("dataset", dataset), ("specie", specie), ("rsa_model", rsa_model)]:
        if key in group_metadata and group_metadata[key] != expected:
            raise ValueError(f"Group metadata disagrees with report: {key}")
    if regression_model is not None:
        for receipt in (correction, group_metadata):
            for key, expected in [("target_model", rsa_model), ("regression_model", regression_model),
                                  ("statistic", "beta")]:
                if key in receipt and receipt[key] != expected:
                    raise ValueError(f"Regression metadata disagrees with report: {key}")

    subjects = set()
    files = group_metadata.get("file_list", [])
    for filename in files:
        match = re.search(rf"(?:^|[\\/]){re.escape(specie)}-sub-(\d+)(?:[\\/]|$)", str(filename))
        if match:
            subjects.add(int(match[1]))
    n_subjects = len(subjects) if files and len(subjects) and all(
        re.search(rf"(?:^|[\\/]){re.escape(specie)}-sub-\d+(?:[\\/]|$)", str(f)) for f in files
    ) else None
    if n_subjects is None:
        warnings.append("Sample size unavailable from actual group input files; omitted from prose.")

    title = " ".join((report_title or rsa_model.replace("_", " ")).split())
    species_label = {"D": "dogs", "H": "humans"}.get(specie, specie)
    subject_clause = f" (n = {n_subjects})" if n_subjects else ""
    intro = f"In {dataset}, searchlight RSA for {title} in {species_label}{subject_clause}"
    if regression_model is not None:
        intro = (f"In {dataset}, regression searchlight RSA for {title} in {species_label}{subject_clause}, "
                 f"controlling for the models specified by {regression_model},")
    comparison = correction.get("threshold_comparison", ">" if regression_model is not None else ">=")
    if comparison not in {">", ">="}:
        raise ValueError("Unsupported saved threshold comparison")
    threshold_symbol = ">" if comparison == ">" else "≥"
    threshold = correction.get("z_threshold")
    minimum = correction.get("minimal_cluster_size")
    if threshold is not None and (not np.isfinite(threshold) or threshold <= 0):
        raise ValueError("Invalid saved z threshold")
    if minimum is not None and (minimum < 1 or int(minimum) != minimum):
        raise ValueError("Invalid saved cluster extent threshold")
    if clusters and threshold is not None and np.any(data[labelled > 0] < threshold - 1e-6):
        raise ValueError("Corrected image contains voxels below saved threshold")
    if clusters and threshold is not None and comparison == ">" and np.any(data[labelled > 0] <= threshold):
        raise ValueError("Corrected image violates the strict saved threshold")
    if clusters and minimum is not None and np.any(sizes < minimum):
        raise ValueError("Corrected image contains clusters below saved extent threshold")
    threshold_text = ""
    if threshold is not None and minimum is not None:
        threshold_text = (f" at a cluster-forming threshold of Z {threshold_symbol} {_number(threshold)}"
                          f" and a minimum cluster extent of {minimum} voxels")
    else:
        warnings.append("Correction thresholds incomplete; omitted from Results prose.")
    if clusters:
        noun = "cluster" if n_clusters == 1 else "clusters"
        text = f"{intro} yielded {n_clusters} surviving {noun}, comprising {n_voxels} voxels{threshold_text}."
        if n_clusters > 1:
            text += f" Cluster sizes ranged from {int(sizes.min())} to {int(sizes.max())} voxels."
        largest = max(clusters, key=lambda c: c["size_vox"])
        strongest = max(clusters, key=lambda c: c["peaks"][0]["z"])

        def describe(cluster):
            peak = cluster["peaks"][0]
            location = peak["region"]
            if location == "Unlabelled":
                location = "an unlabelled atlas location"
            elif location == "Outside atlas":
                location = "a location outside the atlas"
            else:
                location = "the " + location[0].lower() + location[1:]
            coords = ", ".join(_number(v) for v in peak["xyz_mm"])
            return (f"{location} ({cluster['size_vox']} voxels; peak Z = {peak['z']:.2f}; "
                    f"coordinates: {coords} mm)")

        if largest == strongest:
            text += f" The {'cluster' if n_clusters == 1 else 'largest cluster'} contained the highest peak statistic at {describe(largest)}."
        else:
            text += f" The largest cluster had its maximum at {describe(largest)}."
            text += f" The highest peak statistic occurred at {describe(strongest)}."
        secondary = list(dict.fromkeys(p["region"] for p in strongest["peaks"][1:]
                                      if p["region"] not in {"Unlabelled", "Outside atlas"}))
        if secondary:
            names = [name[0].lower() + name[1:] for name in secondary]
            locations = " and ".join(names) if len(names) < 3 else ", ".join(names[:-1]) + ", and " + names[-1]
            text += f" Reported secondary {'peaks' if len(names) > 1 else 'peak'} in this cluster "
            text += f"{'were' if len(names) > 1 else 'was'} labelled as {locations}."
    else:
        text = f"{intro} yielded no surviving clusters{threshold_text}."

    voxel_volume = float(abs(np.linalg.det(image.affine[:3, :3])))
    notes = ["Coordinates are world coordinates in millimetres in the corrected result image's template space.",
             "Extent is reported once per cluster; indented rows are secondary peaks.",
             "Anatomical labels describe peak locations, not the entire cluster.",
             "Z denotes the standardized RSA statistic, not the RSA correlation coefficient.",
             f"Voxel volume: {_number(voxel_volume)} mm³."]
    if regression_model is not None:
        notes[3] = ("Z denotes the group target regression coefficient standardized against its permutation null, "
                    "not the coefficient itself or a regression t statistic. Inference tests the positive tail.")
    if threshold_text:
        notes.append(f"Cluster-forming threshold: Z {threshold_symbol} {_number(threshold)}; minimum extent: {minimum} voxels.")
    alpha = correction.get("cluster_threshold")
    if alpha is not None:
        if not 0 < alpha < 1:
            raise ValueError("Invalid saved cluster probability threshold")
        if correction.get("forced_minimal_cluster_size") is None and (
                "forced_minimal_cluster_size" in correction or
                (regression_model is not None and correction.get("cluster_distribution"))):
            notes.append(f"Saved cluster-correction probability threshold: {_number(alpha)}.")
        else:
            warnings.append("Extent may have been manually forced; no probability claim included.")
    if atlas_name:
        notes.append(f"Peak-label atlas: {atlas_name}.")
    methods = {key: group_metadata[key] for key in ["model", "dis_method", "rsa_method", "radius", "mah_fold", "mask_type"]
               if key in group_metadata}
    if methods.get("dis_method"):
        notes.append(f"Neural-pattern dissimilarity method: {methods['dis_method']}.")
    if regression_model is not None:
        methods.update(regression_model=regression_model, target_model=rsa_model,
                       statistic="beta", threshold_comparison=comparison)
        if group_metadata.get("weighting"):
            methods["weighting"] = group_metadata["weighting"]
            notes.append(f"Group weighting: {group_metadata['weighting']}.")
    if max_peaks_per_cluster is not None:
        notes.append(f"Up to {max_peaks_per_cluster} peaks per cluster, separated by at least {_number(min_dist_mm)} mm.")
    unknown = sum(p["region"] in {"Unlabelled", "Outside atlas"} for c in clusters for p in c["peaks"])
    if unknown:
        notes.append("Unlabelled indicates no named atlas region; outside atlas indicates coordinates outside its grid.")
        warnings.append(f"{unknown} reported peak(s) have no named atlas region.")
    sources = {}
    for name, path in [("csv", csv_path), ("corrected_map", corrected_map),
                       ("correction_metadata", correction_path), ("group_metadata", group_path)]:
        sources[name] = {"path": str(path.resolve()), "sha256": hashlib.sha256(path.read_bytes()).hexdigest() if path.is_file() else None}
    caption = f"Searchlight RSA for {title} in {species_label} ({dataset})."
    if regression_model is not None:
        caption = (f"Regression searchlight RSA for {title} in {species_label} ({dataset}), "
                   f"controlling for the models specified by {regression_model}.")
    return dict(schema_version=1, dataset=dataset, specie=specie, rsa_model=rsa_model,
                regression_model=regression_model, title=title, caption=caption,
                results_text=text, notes=notes, warnings=warnings, clusters=clusters,
                n_clusters=int(n_clusters), n_voxels=n_voxels, n_subjects=n_subjects,
                voxel_volume_mm3=voxel_volume, correction=correction, analysis=methods,
                sources=sources, atlas_name=atlas_name,
                peak_selection=dict(min_dist_mm=min_dist_mm, max_peaks_per_cluster=max_peaks_per_cluster))


def table_rows(report):
    for cluster in report["clusters"]:
        for index, peak in enumerate(cluster["peaks"]):
            yield [str(cluster["cluster_id"]) if index == 0 else "",
                   peak["region"], str(cluster["size_vox"]) if index == 0 else "",
                   f"{peak['z']:.2f}", *[_number(v) for v in peak["xyz_mm"]]]


def _latex(text):
    replacements = {"\\": r"\textbackslash{}", "&": r"\&", "%": r"\%", "$": r"\$",
                    "#": r"\#", "_": r"\_", "{": r"\{", "}": r"\}",
                    "~": r"\textasciitilde{}", "^": r"\textasciicircum{}",
                    "≥": r"\ensuremath{\geq}", ">": r"\ensuremath{>}", "³": r"\textsuperscript{3}"}
    return "".join(replacements.get(char, char) for char in str(text))


def write_latex(report, path):
    identity = [report["dataset"], report["specie"], report["rsa_model"], report.get("regression_model")]
    label = "tab:rsa-" + hashlib.sha256(json.dumps(identity).encode()).hexdigest()[:12]
    lines = [r"% Include with \input{...}; requires \usepackage{booktabs,longtable,array}.",
             _latex(report["results_text"]), "", r"\begingroup", r"\small",
             r"\setlength{\tabcolsep}{3pt}",
             r"\begin{longtable}{@{}r>{\raggedright\arraybackslash}p{0.32\linewidth}rrrrr@{}}",
             r"\caption{" + _latex(report["caption"]) + r"}\label{" + label + r"}\\",
             r"\toprule", " & ".join(_latex(h.replace("\n", " ")) for h in HEADERS) + r" \\",
             r"\midrule", r"\endfirsthead", r"\toprule",
             " & ".join(_latex(h.replace("\n", " ")) for h in HEADERS) + r" \\",
             r"\midrule", r"\endhead", r"\midrule", r"\endfoot", r"\bottomrule", r"\endlastfoot"]
    for row in table_rows(report):
        cells = [_latex(v) for v in row]
        if not row[0]:
            cells[1] = r"\hspace*{1em}" + cells[1]
        lines.append(" & ".join(cells) + r" \\")
    if not report["clusters"]:
        lines.append(r"\multicolumn{7}{l}{No surviving clusters.} \\")
    lines += [r"\end{longtable}", r"\noindent\textit{Note.} " + _latex(" ".join(report["notes"])), r"\endgroup", ""]
    Path(path).write_text("\n".join(lines), encoding="utf-8")


def write_docx(report, path):
    if not word_export_supported():
        print(f"WARNING: {REMOTE_WARNING}")
        return False
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("Step 10 Word export requires python-docx. Install requirements-reporting.txt in the pipeline's Python environment.") from exc
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    section.top_margin = section.bottom_margin = Inches(0.7)
    section.left_margin = section.right_margin = Inches(0.65)
    for name in ["Normal", "Title", "Heading 1", "Caption"]:
        style = document.styles[name]
        style.font.name = "Times New Roman"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.size = Pt(11 if name != "Title" else 15)
        # Templates supplied by different python-docx installations can carry
        # themed fonts and decorative title rules. Make the paper style explicit.
        fonts = style._element.get_or_add_rPr().rFonts
        for key in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn("w:" + key), "Times New Roman")
            fonts.attrib.pop(qn("w:" + key + "Theme"), None)
        for border in style._element.findall(".//" + qn("w:pBdr")):
            border.getparent().remove(border)
    normal = document.styles["Normal"].paragraph_format
    normal.space_after, normal.line_spacing = Pt(7), 1.08
    document.add_paragraph(f"{report['dataset']} {report['title']} RSA results", "Title")
    document.add_paragraph(report["results_text"])
    caption = document.add_paragraph("Table 1. " + report["caption"], "Caption")
    caption.paragraph_format.keep_with_next = True
    table = document.add_table(rows=1, cols=7)
    table.autofit = False
    widths = [0.72, 2.65, 0.72, 0.65, 0.72, 0.72, 0.72]
    for col, width in zip(table.columns, widths):
        col.width = Inches(width)
    for cell, header in zip(table.rows[0].cells, HEADERS):
        cell.text = header
    repeat = OxmlElement("w:tblHeader")
    table.rows[0]._tr.get_or_add_trPr().append(repeat)
    rows = list(table_rows(report))
    for values in rows:
        for cell, value in zip(table.add_row().cells, values):
            cell.text = value
    if not rows:
        table.add_row().cells[0].merge(table.rows[-1].cells[-1]).text = "No surviving clusters."
    # Journal-style horizontal rules intentionally follow the requested paper layout.
    borders = OxmlElement("w:tblBorders")
    for name in ["top", "bottom", "left", "right", "insideH", "insideV"]:
        edge = OxmlElement("w:" + name)
        edge.set(qn("w:val"), "single" if name in {"top", "bottom"} else "nil")
        edge.set(qn("w:sz"), "8")
        borders.append(edge)
    table._tbl.tblPr.append(borders)
    for ri, row in enumerate(table.rows):
        no_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(no_split)
        for ci, cell in enumerate(row.cells):
            if rows or ri == 0:
                cell.width = Inches(widths[ci])
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(3)
                paragraph.paragraph_format.space_after = Pt(3)
                paragraph.paragraph_format.line_spacing = 1
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if ci == 1 else WD_ALIGN_PARAGRAPH.RIGHT
                if ri > 0 and rows and not rows[ri - 1][0] and ci == 1:
                    paragraph.paragraph_format.left_indent = Inches(0.12)
                for run in paragraph.runs:
                    run.font.size = Pt(10)
                    run.bold = ri == 0
            if ri == 0:
                cb = OxmlElement("w:tcBorders")
                edge = OxmlElement("w:bottom")
                edge.set(qn("w:val"), "single")
                edge.set(qn("w:sz"), "6")
                cb.append(edge)
                cell._tc.get_or_add_tcPr().append(cb)
    note = document.add_paragraph("Note. " + " ".join(report["notes"]))
    for run in note.runs:
        run.font.size = Pt(9)
    document.core_properties.title = report["caption"]
    document.core_properties.author = ""
    document.save(path)


def export_publication_report(csv_path, corrected_map, **context):
    """Write portable reports, plus Word only on Windows/Colab; return created paths."""
    allow_word = word_export_supported()
    if not allow_word:
        print(f"WARNING: {REMOTE_WARNING}")
    report = build_report(csv_path, corrected_map, **context)
    stem = str(Path(csv_path).with_suffix(""))
    outputs = [stem + "_publication.docx", stem + "_publication.tex",
               stem + "_results.txt", stem + "_report.json"]
    report["word_export_enabled"] = allow_word
    report["skipped_outputs"] = [] if allow_word else [outputs[0]]
    if allow_word:
        write_docx(report, outputs[0])
    else:
        report["warnings"].append(REMOTE_WARNING)
    write_latex(report, outputs[1])
    Path(outputs[2]).write_text(report["results_text"] + "\n", encoding="utf-8")
    Path(outputs[3]).write_text(json.dumps(report, indent=2, ensure_ascii=False, default=str) + "\n", encoding="utf-8")
    for warning in report["warnings"]:
        if warning != REMOTE_WARNING:
            print(f"Report note: {warning}")
    return outputs if allow_word else outputs[1:]


def mirror_results(file_pairs, *, root=DRIVE_RESULTS_ROOT):
    """Only an existing Windows Drive results root permits any directory creation."""
    if os.name != "nt" or not os.path.isdir(root):
        print("Skipped Google Drive copy: requires Windows and an existing Google Drive Results folder.")
        return False
    try:
        for source, destination in file_pairs:
            os.makedirs(os.path.dirname(destination), exist_ok=True)
            shutil.copyfile(source, destination)
        print("Copied result files to Google Drive.")
        return True
    except OSError as exc:
        print(f"WARNING: skipped Google Drive copy ({exc}); primary outputs are unchanged.")
        return False
